"""Regenerate editable demo assets: python3 scripts/generate-downloads.py.
Requires python-docx and openpyxl. PDF is printed by scripts/generate-pdf.mjs.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
out = Path(__file__).resolve().parents[1] / 'public' / 'downloads'
out.mkdir(parents=True, exist_ok=True)
doc=Document()
style=doc.styles['Normal']; style.font.name='Noto Sans JP'; style.font.size=Pt(11)
style._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans JP')
doc.add_heading('Careety | ES作成テンプレート',0)
doc.add_paragraph('大学生活も、就活も。経験を整理して、自分の言葉で伝えよう。')
doc.add_paragraph('デザイン確認用のサンプル資料です。選考結果を保証するものではありません。')
for title, guide in [ ('1. 結論','何に力を入れたか・どんな強みを伝えるかを一文で。'),('2. 背景・課題','どのような状況で、何を課題と考えたか。'),('3. 自分の行動','何を考え、どう工夫したか。自分の役割を具体的に。'),('4. 結果・学び','結果として何が変わり、何を学んだか。事実に基づいて。'),('5. ESの下書き','設問と文字数を確認して、上のメモを文章にまとめよう。') ]:
 doc.add_heading(title,2);doc.add_paragraph(guide);doc.add_paragraph('\n\n____________________________________________\n')
doc.add_heading('提出前のチェック',2)
for text in ['設問に答えている','結論から書いている','自分の行動が具体的','文字数・企業名・誤字を確認','実際の経験と一致している']:
 doc.add_paragraph('□ '+text)
doc.save(out/'careety-es-template.docx')
def workbook(filename,title,heads,rows,widths):
 wb=Workbook(); ws=wb.active;ws.title='ワークシート';ws.append([title]);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(heads));ws.append(['Careety サンプル資料｜大学生活も、就活も。']);ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(heads));ws.append(heads)
 for row in rows:ws.append(row)
 for i,w in enumerate(widths,1):ws.column_dimensions[chr(64+i)].width=w
 for cell in ws[1]:cell.font=Font(name='Noto Sans JP',size=16,bold=True,color='0158C2')
 for cell in ws[3]:cell.fill=PatternFill('solid',fgColor='0158C2');cell.font=Font(name='Noto Sans JP',bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=32;ws.row_dimensions[3].height=28
 for row in ws.iter_rows(min_row=4):
  ws.row_dimensions[row[0].row].height=65
  for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top');cell.font=Font(name='Noto Sans JP',size=11);cell.fill=PatternFill('solid',fgColor='F7FBFF' if cell.row%2==0 else 'FFFFFF')
 ws.freeze_panes='A4';ws.auto_filter.ref=f'A3:{chr(64+len(heads))}{ws.max_row}';ws.sheet_view.showGridLines=False
 wb.save(out/filename)
workbook('careety-self-analysis.xlsx','自己分析シート',['経験・出来事','状況・課題','自分がしたこと','そのときの気持ち','なぜそう感じた？','見つけた価値観'],[['嬉しかった経験','','','','',''],['難しかった経験','','','','',''],['続けてきた経験','','','','','']]+[['']*6 for _ in range(7)],[24,28,30,28,30,28])
workbook('careety-schedule.xlsx','就活スケジュール表',['日付','大学生活の予定','就活の予定・企業名','締め切り','次の行動','完了チェック'],[['記入例','ゼミ発表','興味のある企業を調べる','公式サイトで確認','応募条件をメモする','未完了']]+[['']*6 for _ in range(15)],[20,28,32,24,35,18])
print('Generated Word and Excel templates.')
